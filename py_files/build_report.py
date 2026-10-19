"""Build a reproducible AVP figure set and an editable report template.

Run from the repository root: python -m py_files.build_report
The calculation follows the low-frequency model documented in info.md.
"""

from pathlib import Path
from copy import deepcopy
import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from lxml import etree
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from py_files import config as cfg
from py_files.functions import (
    get_barlet_values,
    get_cap_values,
    get_channel_matrix,
    get_covariation_matrix,
    get_direction_matrix,
    get_est_phi,
    get_gamma_array,
    get_tau_array,
)


ROOT = Path(__file__).resolve().parents[1]
IMAGE_DIR = ROOT / "plots_image"
TRUE_PHI = np.pi / 4
SEED = 20261019


def report_data():
    params = cfg.CONFIG_PARAMETRS.copy()
    params["M"] = 3
    params["target_phi"] = TRUE_PHI
    params["main_rng"] = np.random.default_rng(SEED)
    gamma = get_gamma_array(params["M"])
    tau = get_tau_array(params)
    channels = get_channel_matrix(params, tau, cfg.t)
    W = get_covariation_matrix(channels, params["N"])
    directions = get_direction_matrix(params["phis"], params["M"])
    barlett = get_barlet_values(W, directions)
    capon = get_cap_values(W, directions)

    assert channels.shape == (9, params["N"])
    assert np.allclose(W, W.T, atol=1e-12)
    assert np.linalg.eigvalsh(W).min() > 0
    assert np.all(np.isfinite(barlett)) and np.all(np.isfinite(capon))
    estimates = {
        "barlett": get_est_phi(params["phis"], barlett),
        "capon": get_est_phi(params["phis"], capon),
    }
    # The repository's real-valued steering vector neglects the small
    # intersensor phase shifts. At alpha=10 its expected bias is limited.
    for estimate in estimates.values():
        angular_error = np.angle(np.exp(1j * (estimate - TRUE_PHI)))
        assert abs(angular_error) < np.deg2rad(3), (estimate, angular_error)
    return params, gamma, tau, channels, barlett, capon, estimates


def set_plot_style():
    plt.rcParams.update({
        "font.family": "Times New Roman",
        "font.size": 14,
        "mathtext.fontset": "stix",
        "axes.titlesize": 16,
        "axes.labelsize": 14,
        "legend.fontsize": 11,
        "xtick.labelsize": 12,
        "ytick.labelsize": 12,
        "savefig.dpi": 220,
    })


def save_figure(fig, name):
    path = IMAGE_DIR / name
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def build_figures(params, gamma, tau, channels, barlett, capon, estimates):
    IMAGE_DIR.mkdir(exist_ok=True)
    set_plot_style()
    outputs = {}

    fig, ax = plt.subplots(figsize=(6.3, 5.0), constrained_layout=True)
    theta = np.linspace(0, 2 * np.pi, 500)
    R = params["R"]
    ax.plot(R * np.cos(theta), R * np.sin(theta), color="blue", lw=2)
    x, y = R * np.cos(gamma), R * np.sin(gamma)
    ax.scatter(x, y, s=85, color="red", zorder=3)
    for index, (xx, yy) in enumerate(zip(x, y), start=1):
        ax.annotate(f"Д{index}", (xx, yy), xytext=(7, 7),
                    textcoords="offset points")
    ax.annotate("", xy=(1.4 * R * np.cos(TRUE_PHI),
                         1.4 * R * np.sin(TRUE_PHI)), xytext=(0, 0),
                arrowprops={"arrowstyle": "->", "color": "green", "lw": 2.5})
    ax.text(0.87 * R, 1.10 * R, r"$\varphi_0=\pi/4$", color="green")
    ax.axhline(0, color="0.65", lw=0.8)
    ax.axvline(0, color="0.65", lw=0.8)
    ax.set(xlabel="x, м", ylabel="y, м", title="Круговая решётка АВП")
    ax.grid(True, alpha=0.3)
    ax.set_aspect("equal", adjustable="box")
    ax.set_xlim(-1.5 * R, 1.65 * R)
    ax.set_ylim(-1.5 * R, 1.65 * R)
    outputs["geometry"] = save_figure(fig, "report_avp_geometry.png")

    t = cfg.t
    ideal = params["A"] * np.sin(2 * np.pi * params["f0"] * (t - tau[0]))
    noisy = channels[0]
    visible = t <= cfg.TIME_LIMIT["right"]
    fig, ax = plt.subplots(figsize=(7.0, 3.5), constrained_layout=True)
    ax.plot(t[visible], noisy[visible], color="blue", lw=0.8,
            label="Входной сигнал с шумом")
    ax.plot(t[visible], ideal[visible], color="red", lw=1.8,
            label="Сигнал без шума")
    ax.set(xlabel="Время, с", ylabel="Амплитуда, усл. ед.",
           title="Входной сигнал первого АВП")
    ax.grid(True, alpha=0.4)
    ax.legend(loc="upper right", framealpha=0.95)
    outputs["signal"] = save_figure(fig, "report_noisy_input.png")

    ticks = np.arange(9) * np.pi / 4
    labels = ["0", r"$\pi/4$", r"$\pi/2$", r"$3\pi/4$", r"$\pi$",
              r"$5\pi/4$", r"$3\pi/2$", r"$7\pi/4$", r"$2\pi$"]
    for key, values, color, title, method in [
        ("barlett", barlett, "red", "Спектр Бартлетта", "Бартлетт"),
        ("capon", capon, "green", "Спектр Кейпона", "Кейпон"),
    ]:
        fig, ax = plt.subplots(figsize=(7.0, 3.8), constrained_layout=True)
        ax.plot(params["phis"], values / np.max(values), color=color,
                lw=2.5, label=f"Метод {method}")
        ax.axvline(TRUE_PHI, color="blue", linestyle="--", lw=1.8,
                   label=r"Истинный угол $\pi/4$")
        peak = np.max(values / np.max(values))
        ax.scatter([estimates[key]], [peak], color="black", s=40, zorder=5,
                   label=rf"Оценка {estimates[key]:.3f} рад")
        ax.set(xlabel=r"Азимут $\varphi$, рад", ylabel="Нормированный спектр",
               title=title, xlim=(0, 2 * np.pi), ylim=(0, 1.08))
        ax.set_xticks(ticks, labels)
        ax.grid(True, alpha=0.4)
        ax.legend(loc="upper right", framealpha=0.95)
        outputs[key] = save_figure(fig, f"report_{key}_spectrum.png")
    return outputs


def set_font(style, size=14, bold=False):
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), "Times New Roman")


def add_text(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def add_formula(doc, mathml, number):
    """Insert an editable Word equation with its number at the right margin."""
    office_xsl = (Path(os.environ["ProgramFiles(x86)"])
                  / "Microsoft Office" / "root" / "Office16" / "MML2OMML.XSL")
    converter = etree.XSLT(etree.parse(str(office_xsl)))
    math = etree.fromstring(
        ('<math xmlns="http://www.w3.org/1998/Math/MathML">'
         + mathml + '</math>').encode("utf-8")
    )
    omml = converter(math).getroot()

    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    widths = (Cm(1.3), Cm(13.9), Cm(1.3))
    for cell, width in zip(table.rows[0].cells, widths):
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    table.rows[0].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.rows[0].cells[1].paragraphs[0]._p.append(deepcopy(omml))
    number_p = table.rows[0].cells[2].paragraphs[0]
    number_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    number_p.add_run(f"({number})")
    table.rows[0].cells[0].paragraphs[0].add_run("")
    for cell in table.rows[0].cells:
        cell.paragraphs[0].paragraph_format.space_before = Pt(3)
        cell.paragraphs[0].paragraph_format.space_after = Pt(3)
        cell.paragraphs[0].paragraph_format.first_line_indent = Cm(0)
        cell.paragraphs[0].paragraph_format.left_indent = Cm(0)
        cell.paragraphs[0].paragraph_format.right_indent = Cm(0)
    return table


def add_figure(doc, path, number, caption, width_cm=15.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    p = doc.add_paragraph(f"Рисунок {number} — {caption}", style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True


def build_document(params, figures):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    normal = doc.styles["Normal"]
    set_font(normal)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.space_after = Pt(0)
    for name in ("Title", "Heading 1", "Heading 2", "Caption"):
        set_font(doc.styles[name], 14, name != "Caption")
        doc.styles[name].paragraph_format.space_before = Pt(10)
        doc.styles[name].paragraph_format.space_after = Pt(8)
        doc.styles[name].paragraph_format.keep_with_next = name != "Caption"
    doc.styles["Title"].paragraph_format.space_before = Pt(0)
    doc.styles["Title"].paragraph_format.space_after = Pt(18)
    title_ppr = doc.styles["Title"].element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    footer._p.append(page_field)

    title = doc.add_paragraph("Пеленгование источника звука круговой решёткой АВП", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("1 Краткие теоретические сведения", 1)
    add_text(doc, "Три акустических векторных приёмника (АВП) расположены равномерно "
             "на окружности радиуса R. Каждый АВП имеет канал давления y₀ и два "
             "ортогональных векторных канала y₁, y₂. Азимут источника отсчитывается "
             "от положительного направления оси x против часовой стрелки.")
    add_formula(doc, "<msub><mi>γ</mi><mi>m</mi></msub><mo>=</mo>"
                "<mfrac><mrow><mn>2</mn><mi>π</mi><mo>(</mo><mi>m</mi>"
                "<mo>−</mo><mn>1</mn><mo>)</mo></mrow><mi>M</mi></mfrac>", 1)
    add_formula(doc, "<msub><mi>τ</mi><mi>m</mi></msub><mo>=</mo>"
                "<mfrac><mi>R</mi><mi>c</mi></mfrac><mi>cos</mi><mo>(</mo>"
                "<msub><mi>φ</mi><mn>0</mn></msub><mo>−</mo>"
                "<msub><mi>γ</mi><mi>m</mi></msub><mo>)</mo>", 2)
    add_text(doc, "В модели сигнал на каждом приёмнике представляет собой гармонику "
             "с независимым белым гауссовским шумом "
             "в каждом канале. Вектор наблюдений содержит 3M каналов. При "
             "низкочастотном приближении используется направляющий вектор, "
             "составленный из повторяющихся троек [1, cos φ, sin φ].")
    add_formula(doc, "<mi>P</mi><mo>(</mo><mi>t</mi><mo>−</mo>"
                "<msub><mi>τ</mi><mi>m</mi></msub><mo>)</mo><mo>=</mo>"
                "<mi>A</mi><mi>sin</mi><mo>(</mo><mn>2</mn><mi>π</mi>"
                "<mi>f</mi><mo>(</mo><mi>t</mi><mo>−</mo>"
                "<msub><mi>τ</mi><mi>m</mi></msub><mo>)</mo><mo>)</mo>", 3)
    add_formula(doc, "<mi>a</mi><mo>(</mo><mi>φ</mi><mo>)</mo><mo>=</mo>"
                "<msup><mrow><mo>[</mo><mn>1</mn><mo>,</mo><mi>cos</mi>"
                "<mi>φ</mi><mo>,</mo><mi>sin</mi><mi>φ</mi><mo>,</mo>"
                "<mi>…</mi><mo>]</mo></mrow><mi>T</mi></msup>", 4)
    add_formula(doc, "<mi>W</mi><mo>=</mo><mfrac><mrow><mi>X</mi>"
                "<msup><mi>X</mi><mi>T</mi></msup></mrow><mi>N</mi></mfrac>", 5)
    add_formula(doc, "<msub><mi>E</mi><mi>Б</mi></msub><mo>(</mo><mi>φ</mi>"
                "<mo>)</mo><mo>=</mo><msup><mi>a</mi><mi>T</mi></msup>"
                "<mo>(</mo><mi>φ</mi><mo>)</mo><mi>W</mi><mi>a</mi>"
                "<mo>(</mo><mi>φ</mi><mo>)</mo>", 6)
    add_formula(doc, "<msub><mi>E</mi><mi>К</mi></msub><mo>(</mo><mi>φ</mi>"
                "<mo>)</mo><mo>=</mo><mfrac><mn>1</mn><mrow>"
                "<msup><mi>a</mi><mi>T</mi></msup><mo>(</mo><mi>φ</mi>"
                "<mo>)</mo><msup><mi>W</mi><mrow><mo>−</mo><mn>1</mn>"
                "</mrow></msup><mi>a</mi><mo>(</mo><mi>φ</mi>"
                "<mo>)</mo></mrow></mfrac>", 7)
    add_text(doc, "Оценка азимута для каждого метода определяется положением максимума "
             "спектра. На рисунках спектры разделены на собственный максимум; "
             "это меняет только масштаб по вертикали. Отношение С/Ш в коде "
             "определено как A²/σ², где σ — среднеквадратическое отклонение шума.")

    doc.add_heading("2 Параметры моделирования", 1)
    add_text(doc, f"M = {params['M']}; R = {params['R']:.4f} м; c = {params['c']:.2f} м/с; "
             f"f = {params['f0']} Гц; A = {params['A']}; "
             f"φ₀ = π/4 рад (45°); частота дискретизации {cfg.fs} Гц; "
             f"N = {params['N']}; С/Ш = {params['snr']}; "
             f"λ/(2R) = {params['alpha']}. Для воспроизводимости задано "
             f"зерно генератора {SEED}.")
    add_text(doc, "Геометрия круговой решётки и истинный азимут показаны на рисунке 1.")
    add_figure(doc, figures["geometry"], 1,
               "Расположение трёх АВП и направление на источник", 10.5)

    doc.add_heading("3 Результаты моделирования", 1)
    add_text(doc, "На рисунке 2 показан зашумлённый входной сигнал канала давления "
             "первого АВП в сравнении с сигналом без шума.")
    add_figure(doc, figures["signal"], 2,
               "Входной сигнал первого АВП с шумом и без шума")
    add_text(doc, "На рисунке 3 приведён нормированный энергетический спектр "
             "Бартлетта. Пунктирная линия обозначает истинный азимут, "
             "чёрная точка — положение максимума спектра.")
    add_figure(doc, figures["barlett"], 3,
               "Нормированный спектр Бартлетта")
    add_text(doc, "На рисунке 4 приведён нормированный энергетический спектр "
             "Кейпона с теми же обозначениями углов.")
    add_figure(doc, figures["capon"], 4,
               "Нормированный спектр Кейпона")

    doc.add_heading("4 Дисперсия ошибки пеленга", 1)
    add_text(doc, "На рисунках 5–10 показана выборочная дисперсия ошибки оценки "
             "азимута методами Бартлетта и Кейпона по 250 реализациям шума "
             "для каждого значения изменяемого параметра.")
    variance_figures = [
        ("variance_azimuth.png", "истинного азимута",
         "Дисперсия ошибки оценки в зависимости от истинного азимута"),
        ("variance_snr.png", "отношения сигнал/шум",
         "Дисперсия ошибки оценки в зависимости от отношения сигнал/шум"),
        ("variance_radius.png", "радиуса решётки",
         "Дисперсия ошибки оценки в зависимости от радиуса решётки"),
        ("variance_sensors.png", "числа датчиков",
         "Дисперсия ошибки оценки в зависимости от числа датчиков"),
        ("variance_alpha_1_10.png", "параметра α в диапазоне от 1 до 10",
         "Дисперсия ошибки оценки при изменении α от 1 до 10"),
        ("variance_alpha_0_1_1.png", "параметра α в диапазоне от 0,1 до 1",
         "Дисперсия ошибки оценки при изменении α от 0,1 до 1"),
    ]
    for number, (filename, parameter, caption) in enumerate(variance_figures, start=5):
        p = add_text(doc, f"Зависимость дисперсии от {parameter} приведена на рисунке {number}.")
        p.paragraph_format.keep_with_next = True
        add_figure(doc, IMAGE_DIR / "updated" / filename, number, caption, 15.5)

    doc.add_heading("5 Выводы", 1)
    doc.add_paragraph(" ")
    doc.add_paragraph(" ")
    doc.add_paragraph(" ")

    output = ROOT / "report_avp_template_v6.docx"
    doc.save(output)
    return output


if __name__ == "__main__":
    data = report_data()
    figures = build_figures(*data)
    document = build_document(data[0], figures)
    print(f"Report: {document}")
    for name, path in figures.items():
        print(f"{name}: {path}")
    print("Estimates (rad):", data[-1])
