"""Create a Word report in the supplied supplement's style and a Markdown copy."""
from pathlib import Path
import json
import math
import copy
import numpy as np
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from latex2mathml.converter import convert

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
r = json.loads((OUT/'results.json').read_text(encoding='utf-8'))
c = r['checks']
doc = Document(ROOT/'reports/dox/info_snr_supplement.docx')
body = doc._element.body
for child in list(body):
    if child.tag != qn('w:sectPr'):
        body.remove(child)
section = doc.sections[0]
section.page_width=Cm(21); section.page_height=Cm(29.7)
section.left_margin=Cm(3); section.right_margin=Cm(1.5)
section.top_margin=Cm(2); section.bottom_margin=Cm(2)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    style=doc.styles[name]
    style.font.name='Times New Roman'; style.font.size=Pt(14)
    style.font.color.rgb=RGBColor(0,0,0)
    style.paragraph_format.line_spacing=1.5
    style.paragraph_format.space_after=Pt(0)
normal=doc.styles['Normal'].paragraph_format
normal.first_line_indent=Cm(1.25)
normal.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name in ['Heading 1','Heading 2']:
    doc.styles[name].font.bold=True
    doc.styles[name].paragraph_format.first_line_indent=Cm(0)
    doc.styles[name].paragraph_format.space_before=Pt(12)
    doc.styles[name].paragraph_format.keep_with_next=True

xsl=Path('C:/Program Files/Microsoft Office/Office16/MML2OMML.XSL')
transform=etree.XSLT(etree.parse(str(xsl)))
md=[]; eqno=0; figno=0
def fmt(v,n=6): return f'{v:.{n}f}'.replace('.',',')
def title(text):
    p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Cm(0)
    p.add_run(text).bold=True
    md.append('# '+text+'\n')
def h(text):
    doc.add_heading(text,level=2);md.append('## '+text+'\n')
def p(text):
    pp=doc.add_paragraph(text)
    pp.paragraph_format.widow_control=True
    md.append(text+'\n')
def eq(tex):
    global eqno
    eqno+=1
    pp=doc.add_paragraph()
    pp.paragraph_format.first_line_indent=Cm(0)
    pp.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    pp.paragraph_format.line_spacing=1.15
    pp.paragraph_format.space_before=Pt(5);pp.paragraph_format.space_after=Pt(5)
    pp.paragraph_format.keep_together=True
    tabs=pp.paragraph_format.tab_stops
    tabs.add_tab_stop(Cm(8.25),WD_TAB_ALIGNMENT.CENTER)
    tabs.add_tab_stop(Cm(16.45),WD_TAB_ALIGNMENT.RIGHT)
    pp.add_run('\t')
    mathml=etree.fromstring(convert(tex).encode())
    omml=transform(mathml).getroot()
    for mr in omml.iter(qn('m:r')):
        pr=OxmlElement('w:rPr')
        fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Cambria Math');fonts.set(qn('w:hAnsi'),'Cambria Math');pr.append(fonts)
        size=OxmlElement('w:sz');size.set(qn('w:val'),'24');pr.append(size)
        mr.insert(0,pr)
    pp._p.append(omml)
    pp.add_run(f'\t({eqno})').font.size=Pt(12)
    md.append(f'$$\n{tex} \\qquad ({eqno})\n$$\n')
def figure(name,caption):
    global figno
    figno+=1
    pp=doc.add_paragraph();pp.paragraph_format.first_line_indent=Cm(0)
    pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.keep_with_next=True
    pp.add_run().add_picture(str(OUT/'figures'/name),width=Cm(16.4))
    pp=doc.add_paragraph(f'Рисунок {figno} — {caption}',style='Caption')
    pp.paragraph_format.first_line_indent=Cm(0);pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pp.paragraph_format.line_spacing=1.0;pp.paragraph_format.space_after=Pt(6)
    pp.paragraph_format.keep_together=True
    for run in pp.runs:run.font.size=Pt(12);run.italic=False
    md.append(f'![Рисунок {figno} — {caption}](figures/{name})\n')
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1, cols=len(headers));t.style='Table Grid'
    for cell,text in zip(t.rows[0].cells,headers):cell.text=str(text)
    for row in rows:
        for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
    for idx,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if idx==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,cell in enumerate(row.cells):
            if widths:cell.width=Cm(widths[j])
            for pp in cell.paragraphs:
                pp.paragraph_format.first_line_indent=Cm(0);pp.paragraph_format.line_spacing=1
                pp.paragraph_format.space_after=Pt(4);pp.paragraph_format.space_before=Pt(4)
                pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for rr in pp.runs:rr.font.name='Times New Roman';rr.font.size=Pt(11);rr.bold=(idx==0)
    md.append('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+''.join('| '+' | '.join(map(str,row))+' |\n' for row in rows)+'\n')

title('ИССЛЕДОВАНИЕ ИМПУЛЬСА ЛОРЕНЦА ПРИ ОЦЕНИВАНИИ АЗИМУТА КРУГОВОЙ РЕШЁТКОЙ АВП')
pp=doc.add_paragraph('Определение SNR, временная форма, спектр и сравнение с гармоническим сигналом')
pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.first_line_indent=Cm(0)
p('Цель работы — исследовать пеленгование источника, излучающего импульс Лоренца, методами Бартлетта и Кейпона и сопоставить результаты с гармоническим сигналом при одинаковом энергетическом отношении сигнал/шум. Теоретические выражения согласованы с конечным интервалом наблюдения и фактической реализацией расчёта.')
p('Оформление принято по документу info_snr_supplement.docx. Графики и оценки получены из расчётной схемы notebook.ipynb и report.ipynb с использованием исходных функций проекта. В отличие от числового примера дополнения, где τ = 0,01 с, здесь использовано текущее значение duration_pulse = 0,05 с. Значения из прежнего отчёта для синусоиды не переносились: сравнение выполнено заново с текущим определением SNR.')
h('Исходные параметры исследования')
table(['Параметр','Значение'],[
    ('Число АВП M; радиус R','3; 1,7163 м'),('Скорость звука c','343,26 м/с'),
    ('Частота дискретизации fₛ; число отсчётов N','8800 Гц; 10000'),
    ('Интервал наблюдения T; центр tс','1,13625 с; 0,568125 с'),
    ('Ширина импульса τ; коэффициент A','0,05 с; 1 (параметр формулы)'),
    ('Гармонический сигнал','Aₛ = 1 Па; f₀ = 10 Гц'),
    ('Истинный азимут φ₀','π/4 рад = 45°'),
    ('Базовый SNR; исследуемые SNR','2; 0,1…1 с шагом 0,1 и точка 2'),
    ('Реализации шума K; seed','100 на сочетание параметров; 10'),
    ('Угловая сетка','6283 точки; шаг ≈ 0,00100003 рад')], [7.1,9.3])

h('1 Определение SNR и выбор интервала интегрирования')
p('Программа обрабатывает конечную запись, поэтому энергия полезного сигнала и средняя энергия шума должны относиться к одному и тому же окну наблюдения. Временная сетка начинается с нуля; её последний отсчёт определяет правую границу T.')
eq(r'\Delta t=1/f_s,\quad t_n=n\Delta t,\quad n=0,\ldots,N-1')
eq(r'T=(N-1)/f_s=1.13625\ \mathrm{s}')
p('Интегрирование от 0 до T следует из границ доступной записи, а не из предположения, что импульс вне окна равен нулю. Импульс Лоренца имеет бесконечные хвосты. Интеграл по всей временной оси характеризует полную энергию его математической формы, но не энергию реально обработанного фрагмента.')
eq(r'E_s(T)=\int_0^T P^2(t)\,dt')
eq(r'\mathrm{SNR}_T=\frac{E_s(T)}{\mathbb{E}[E_n(T)]}=\frac{E_s(T)}{T\sigma^2}')
p('Шум в реализации — независимые гауссовские отсчёты с нулевым средним и дисперсией σ². Здесь не используется идеальный непрерывный белый шум с неограниченной полосой и бесконечной точечной дисперсией. Энергия конкретной шумовой реализации случайна; заданный SNR определяется через математическое ожидание шумовой энергии.')
eq(r'\sigma^2=\frac{E_s(T)}{T\,\mathrm{SNR}_T},\qquad \sigma=\sqrt{\frac{E_s(T)}{T\,\mathrm{SNR}_T}}')
p('Именно последнее выражение реализовано в calculate_noise_level. Энергия в calculate_energy рассчитывается методом трапеций:')
eq(r'E_s^{(d)}=\Delta t\left[\frac{p_0^2}{2}+\sum_{n=1}^{N-2}p_n^2+\frac{p_{N-1}^2}{2}\right]')
eq(r'\mathbb{E}[E_n^{(d)}]=\Delta t\,(N-1)\sigma^2=T\sigma^2')
p('Последнее равенство точно для принятой дискретной модели, поскольку сумма весов трапеций равна N − 1. Деление энергий на одинаковое T даёт эквивалентное отношение средних мощностей. Пиковое отношение и выражение A²/σ² не совпадают с этим SNR. Логарифмическая запись имеет вид SNRдБ = 10 log₁₀(SNR); базовое значение 2 соответствует 3,0103 дБ.')
p('Шумовая дисперсия каждого АВП определяется по энергии его канала давления. В векторных каналах полезные сигналы дополнительно умножаются на cos φ₀ и sin φ₀. Поэтому их энергетические SNR равны SNR·cos²φ₀ и SNR·sin²φ₀; при φ₀ = 45° они составляют половину SNR канала давления.')

h('2 Формула импульса с учётом положения центра')
eq(r'P(t)=\frac{A\tau}{\tau^2+(t-t_c)^2},\qquad \tau>0')
eq(r'P(t_c)=\frac{A}{\tau},\qquad P(t_c\pm\tau)=\frac{A}{2\tau}')
p('Параметр τ задаёт половину полной ширины на половине максимума; FWHM = 2τ = 0,1 с. Он не является конечной длительностью импульса. В данной формуле A — коэффициент масштаба, а не высота пика: при A = 1 и τ = 0,05 с максимум равен 20 Па. Если P измеряется в Па, размерность коэффициента A здесь Па·с. В формуле синусоиды аналогичный числовой параметр задаёт амплитуду в Па; для строгого различения ниже он обозначен Aₛ.')
p('Для круговой решётки углы расположения АВП и задержки имеют вид:')
eq(r'\gamma_m=\frac{2\pi(m-1)}{M},\qquad d_m=\frac{R}{c}\cos(\varphi_0-\gamma_m)')
eq(r'P_m(t)=P(t-d_m)=\frac{A\tau}{\tau^2+(t-t_c-d_m)^2}')
p('Обозначение dₘ отделяет задержку распространения от параметра ширины τ. В коде dₘ называется tau_array. Центр импульса в канале m находится в tс + dₘ. Для трёх АВП задержки составляют 3,535534; 1,294095; −4,829629 мс соответственно.')

h('3 Временная форма и обоснование центра T/2')
p('Опорный импульс помещён в середину окна: tс = T/2 = 0,568125 с. Это обеспечивает равный запас до обеих границ и максимальную долю зарегистрированной энергии для фиксированных T и τ.')
figure('01_pulse.png','Опорный импульс Лоренца на интервале [0; T]')
p('Утверждение о максимуме можно доказать непосредственно. Обозначим квадрат центрированного импульса через g; g чётна и строго убывает с увеличением положительного аргумента.')
eq(r'g(x)=\frac{A^2\tau^2}{(\tau^2+x^2)^2}')
eq(r'\frac{d E_s(T;t_c)}{dt_c}=g(t_c)-g(T-t_c)')
p('При tс < T/2 расстояние до левой границы меньше расстояния до правой, поэтому производная положительна; при tс > T/2 она отрицательна. Следовательно, единственный максимум достигается при tс = T/2. Выбор центра в нуле оставил бы в окне только правую часть формы и существенно уменьшил зарегистрированную энергию. Этот вывод относится к опорному импульсу; центры отдельных каналов дополнительно сдвигаются задержками.')
p('Для произвольного положения центра удобно использовать первообразную квадрата импульса:')
eq(r'H(x)=\frac{A^2}{2}\left[\frac{x}{x^2+\tau^2}+\frac{1}{\tau}\arctan\frac{x}{\tau}\right]')
eq(r'E_s(T;t_c)=H(T-t_c)-H(-t_c),\qquad E_\infty=\frac{\pi A^2}{2\tau}')
eq(r'\nu(T)=\frac{2}{\pi}\left[\arctan\frac{T}{2\tau}+\frac{\tau T/2}{(T/2)^2+\tau^2}\right]')
p(f'При заданных параметрах E∞ = {fmt(c["E_infinite"])} Па²·с, Eₛ(T) = {fmt(c["E_imp_numeric"])} Па²·с, ν = {fmt(c["nu"],9)}. В окне содержится {fmt(100*c["nu"],6)} % полной энергии. Относительное расхождение аналитической и численной энергий равно {c["energy_relative_error"]:.3e}. Численно рассчитанная оконная энергия уже учитывает хвосты за границами; повторное умножение на ν недопустимо.')
eq(r'E_{s,m}=H(T-t_c-d_m)-H(-t_c-d_m)')
p('В программе этот же результат получается интегрированием P²(u) по сдвинутому вектору u = t − dₘ, то есть на [−dₘ; T − dₘ]. Для текущих задержек энергии трёх каналов давления равны 31,40691924; 31,40692102; 31,40691745 Па²·с. Различие мало благодаря расположению импульса вдали от границ.')
figure('04_sensors.png','Импульсы в каналах давления трёх АВП с учётом задержек')
p(f'При базовом SNR = 2 среднеквадратический уровень шума для опорного импульса составляет σ = {fmt(c["sigma_imp"])} Па. На рисунке ниже использован отдельный демонстрационный seed = 42, сохранённый из исходной ячейки notebook.ipynb. Статистические эксперименты далее используют seed = 10.')
figure('03_noise.png','Импульс без шума и с гауссовским шумом при энергетическом SNR = 2')

h('4 Частотный спектр импульса')
p('Чтобы исключить неоднозначность коэффициентов нормировки, определим прямое преобразование Фурье с частотой f в герцах:')
eq(r'\widetilde P(f)=\int_{-\infty}^{\infty}P(t)e^{-i2\pi ft}\,dt')
p('После замены x = t − tс используется интеграл от τ/(τ² + x²): при замыкании контура в соответствующей полуплоскости его полюсы x = ±iτ дают множитель π exp(−2πτ|f|). Поэтому для полной временной оси получаем:')
eq(r'\widetilde P_\infty(f)=\pi A e^{-2\pi\tau|f|}e^{-i2\pi ft_c}')
eq(r'|\widetilde P_\infty(f)|=\pi |A|e^{-2\pi\tau|f|}')
p('Сдвиг центра меняет фазу на −2πftс, но сохраняет модуль спектра полной формы. Модуль максимален при f = 0 и экспоненциально убывает с ростом |f|; широкой полосе соответствует меньшая τ. Энергетическая спектральная плотность равна |P̃∞(f)|², тогда как исходный ноутбук строит модуль |P̃(f)|. Это разные величины.')
eq(r'\int_{-\infty}^{\infty}|\widetilde P_\infty(f)|^2df=\frac{\pi A^2}{2\tau}=E_\infty')
p('Равенство Парсеваля подтверждает согласованность спектральной нормировки и полной энергии. Однако программа наблюдает только прямоугольное окно [0; T]. Умножение на окно во времени соответствует свёртке спектров в частотной области:')
eq(r'\widetilde P_T(f)=\widetilde P_\infty(f)*W_T(f)')
eq(r'W_T(f)=T e^{-i\pi fT}\operatorname{sinc}(fT),\quad \operatorname{sinc}(x)=\frac{\sin(\pi x)}{\pi x}')
p('Функция calculate_freq_spectr выполняет FFT, делит результат на fₛ и применяет fftshift с последующим взятием модуля. Деление на fₛ эквивалентно множителю Δt и приближает интегральное преобразование, поэтому единица вертикальной оси — Па·с. Шаг частот ДПФ равен fₛ/N = 0,88 Гц, а не 1/T; длительность T здесь определяется крайними отсчётами, тогда как период записи ДПФ равен N/fₛ.')
eq(r'\widetilde P_k^{(d)}=\Delta t\sum_{n=0}^{N-1}p_n e^{-i2\pi kn/N}')
figure('08_spectrum_comparison.png','Модуль спектра конечной записи и аналитический спектр полной формы; справа логарифмическая шкала')
p(f'В нуле аналитический спектр полной формы равен πA = {fmt(math.pi)} Па·с. Для конечного окна интеграл P(t) равен 2A arctan[T/(2τ)] = {fmt(c["area_exact"])} Па·с; ДПФ даёт {fmt(c["fft_zero"])} Па·с. Их небольшое различие обусловлено прямоугольными весами ДПФ вместо трапеций. Отклонение от πA существенно больше, поскольку площадь хвостов P(t) убывает медленнее, чем энергия хвостов P²(t). Поэтому близость ν к единице не означает точного совпадения модулей спектров во всех частотных точках.')
p('Для низкочастотного приближения важно учитывать отсутствие строгой верхней границы спектра импульса. Вместо неё можно ввести симметричную полосу, содержащую 99 % энергии полной формы:')
eq(r'\frac{\int_{-F}^{F}|\widetilde P_\infty(f)|^2df}{E_\infty}=1-e^{-4\pi\tau F}')
p(f'Отсюда F₉₉ = ln(100)/(4πτ) = {fmt(c["f_99_energy"])} Гц. Для этой границы c/(2RF₉₉) = {fmt(c["low_frequency_alpha_at_f99"],3)}, то есть превышает 10. Это энергетическое обоснование принятого приближения; спектральный хвост за F₉₉ при этом сохраняется. Для синусоиды f₀ = 10 Гц параметр c/(2Rf₀) точно равен 10.')

h('5 Оценивание азимута и сравнение с синусоидой')
p('Каналы АВП формируются по модели y₀⁽ᵐ⁾ = Pₘ + e₀⁽ᵐ⁾, y₁⁽ᵐ⁾ = Pₘ cos φ₀ + e₁⁽ᵐ⁾, y₂⁽ᵐ⁾ = Pₘ sin φ₀ + e₂⁽ᵐ⁾. Шумы независимы между датчиками, каналами и реализациями. Матрица X содержит 3M строк каналов и N столбцов отсчётов.')
eq(r'W=\frac{XX^{\mathsf T}}{N},\qquad a(\varphi)=[1,\cos\varphi,\sin\varphi,\ldots]^{\mathsf T}')
p('Матрица W в исходном коде является матрицей вторых моментов без вычитания временного среднего. Для импульса с ненулевым средним это существенно: замена на центрированную ковариацию изменила бы постановку эксперимента. В направляющем векторе задержки не учитываются — применяется низкочастотное приближение исходной модели.')
eq(r'B(\varphi)=a^{\mathsf T}(\varphi)Wa(\varphi)')
eq(r'C(\varphi)=\frac{1}{a^{\mathsf T}(\varphi)W^{-1}a(\varphi)}')
eq(r'\widehat\varphi=\underset{\varphi\in\Phi}{\operatorname{argmax}}\,B(\varphi)\quad\text{или}\quad\underset{\varphi\in\Phi}{\operatorname{argmax}}\,C(\varphi)')
p('В реализации используются обычная обратная матрица для Кейпона и дискретный поиск максимума. Диагональная нагрузка и интерполяция максимума не добавлялись. Половина шага угловой сетки составляет около 0,02865°; это масштаб квантования, но не гарантированная граница полной ошибки при наличии шума. Метод MUSIC не включён, поскольку в текущих расчётных функциях и ноутбуках он не реализован.')
p('Сигнал сравнения имеет вид Pₛ(t) = Aₛ sin(2πf₀t). Для обоих сигналов оставлены одинаковые окно, решётка, угловая сетка и целевой энергетический SNR; шум масштабируется отдельно по фактической энергии каждого сигнала. Это сравнение при равном SNR, а не при одинаковой пиковой амплитуде, одинаковой энергии или одной общей абсолютной дисперсии шума.')
table(['Величина','Импульс Лоренца','Синусоида'],[
    ('Энергия в окне, Па²·с',fmt(c['E_imp_numeric']),fmt(c['E_sin'])),
    ('Максимальная амплитуда, Па','20','1'),
    ('СКО шума при SNR = 2, Па',fmt(c['sigma_imp']),fmt(c['sigma_sin'])),
    ('Центр / частота','tс = 0,568125 с','f₀ = 10 Гц')])
h('5.1 Оценки для одной реализации')
p('В этой проверке для каждого типа сигнала заново создаётся генератор с seed = 10. Оба метода обрабатывают одну матрицу наблюдений данного сигнала. Стандартизованные шумы между типами сигналов одинаковы, но их абсолютный масштаб различается вследствие энергетической нормировки.')
figure('09_direction_comparison.png','Нормированные диаграммы Бартлетта и Кейпона при SNR = 2')
table(['Сигнал','Метод','Оценка, град.','Ошибка, град.'],[(d['signal'],d['method'],fmt(d['estimate_deg'],6),fmt(d['error_deg'],6)) for d in r['single']])
p('Диаграммы разделены на собственные максимумы; это сохраняет положения оценок. Более узкий максимум Кейпона сам по себе не доказывает меньшую дисперсию ошибки. Совпадение оценок двух методов в этой реализации означает совпадение выбранного узла сетки, а не тождественность алгоритмов. Одной реализации недостаточно для вывода о сравнительной точности.')
h('5.2 Статистический эксперимент')
p('Для каждого типа сигнала, метода и значения SNR выполнено K = 100 реализаций. В каждой паре «сигнал — метод» генератор независимо инициализируется seed = 10 и последовательно расходуется по SNR в возрастающем порядке. Таким образом, сравниваемые методы используют одинаковые шумовые последовательности; объекты конфигурации и генераторы не разделяются между методами. Это устраняет зависимость сравнения от порядка вызова restart_rng в исходном report.ipynb.')
p('Ошибка приводится к главному круговому интервалу, после чего рассчитываются средняя ошибка, несмещённая выборочная дисперсия и корень средней квадратичной ошибки:')
eq(r'\varepsilon_k=((\widehat\varphi_k-\varphi_0+\pi)\bmod 2\pi)-\pi')
eq(r'\overline\varepsilon=\frac1K\sum_{k=1}^{K}\varepsilon_k,\qquad s_\varepsilon^2=\frac{\sum_{k=1}^{K}(\varepsilon_k-\overline\varepsilon)^2}{K-1}')
eq(r'\mathrm{RMSE}=\sqrt{\frac1K\sum_{k=1}^{K}\varepsilon_k^2}')
p('Различие дисперсии и средней квадратичной ошибки учитывается явно: RMSE² = [(K − 1)/K]s² + ε̄². Для наблюдаемых ошибок вблизи нуля обычные выборочные моменты круговой ошибки содержательно интерпретируемы; результаты не относятся к режиму ошибок, распределённых по всей окружности.')
figure('10_statistics_comparison.png','Дисперсия круговой ошибки и RMSE в зависимости от энергетического SNR')
p('Ниже приведены опорные точки; полная таблица по всем 11 значениям SNR сохранена в statistics.csv. Дисперсия указана в радианах в квадрате, RMSE — в градусах.')
subset=[d for d in r['statistics'] if d['snr'] in [0.1,0.5,1,2]]
subset.sort(key=lambda d:(d['snr'],d['signal'],d['method']))
table(['SNR','Сигнал','Метод','s², рад²','RMSE, град.'],[(fmt(d['snr'],1),d['signal'],d['method'],f'{d["variance_rad2"]:.4e}',fmt(math.degrees(d['rmse_rad']),4)) for d in subset])
p('При увеличении SNR общий уровень ошибки снижается, однако отдельные точки не обязаны быть монотонными из-за конечного числа реализаций. В показанной выборке оценки Бартлетта и Кейпона имеют близкую точность. При SNR = 2 RMSE импульса составляет 0,2400° и 0,2403°, а синусоиды — 0,2558° и 0,2577° соответственно. При SNR = 0,5 импульс в этой выборке немного уступает синусоиде; универсального преимущества формы сигнала результаты не устанавливают.')

errors=json.loads((OUT/'errors.json').read_text(encoding='utf-8'))
brng=np.random.default_rng(20261023)
indices=brng.integers(0,100,size=(20000,100))
intervals=[]
for method in ['Бартлетт','Кейпон']:
    a=np.asarray(errors[f'Лоренц|{method}|2']);b=np.asarray(errors[f'Синусоида|{method}|2'])
    ratio=np.var(a,ddof=1)/np.var(b,ddof=1)
    boot=np.var(a[indices],axis=1,ddof=1)/np.var(b[indices],axis=1,ddof=1)
    lo,hi=np.quantile(boot,[.025,.975])
    intervals.append(dict(method=method,variance_ratio=float(ratio),lower=float(lo),upper=float(hi)))
(OUT/'comparison_intervals.json').write_text(json.dumps(intervals,ensure_ascii=False,indent=2),encoding='utf-8')
p('Чтобы оценить неопределённость сравнения в базовой точке SNR = 2, рассчитаны парные процентильные bootstrap-интервалы для отношения дисперсий «импульс / синусоида»: 20000 повторных выборок индексов реализаций с возвращением, отдельный seed = 20261023. Один и тот же набор индексов используется для обоих сигналов, поскольку стандартизованные шумы сопоставлены попарно. Интервалы приближённые и относятся к данной точке, без поправки на множественные сравнения по всей сетке SNR.')
table(['Метод','Отношение s²','Приближённый 95 % интервал'],[(d['method'],fmt(d['variance_ratio'],3),f'[{fmt(d["lower"],3)}; {fmt(d["upper"],3)}]') for d in intervals])
p('Оба интервала включают единицу. Поэтому небольшое наблюдаемое различие при SNR = 2 не является достаточным основанием утверждать превосходство импульса над синусоидой по дисперсии. Неопределённость конечной выборки следует отличать от точности вычисления отдельных оценок на заданной угловой сетке.')

h('6 Выводы')
p('1. Для импульсного сигнала SNR согласован с наблюдаемым окном [0; T] и определяется как Eₛ(T)/(Tσ²). Конечные пределы следуют из доступной записи. Принятый расчёт трапециями обеспечивает точное равенство средней дискретной энергии шума Tσ².')
p('2. Положение tс = T/2 максимизирует зарегистрированную энергию опорного импульса и симметрично располагает его хвосты относительно границ окна. При τ = 0,05 с в запись попадает 99,971335 % полной энергии. Пики канальных сигналов находятся в T/2 + dₘ.')
p('3. При выбранной нормировке преобразования Фурье модуль спектра полной формы равен π|A|exp(−2πτ|f|). Спектр конечной записи отличается из-за прямоугольного окна; доля энергии ν не определяет относительную ошибку модуля спектра в каждой частотной точке.')
p('4. В исследованной конфигурации оба метода восстанавливают азимут вблизи 45°. Для 100 реализаций при SNR = 2 RMSE находится около 0,24° у импульса и 0,26° у синусоиды. Парные bootstrap-интервалы отношения дисперсий включают единицу; устойчивое преимущество одного сигнала или метода по этим данным не установлено.')
p('5. Выводы относятся к текущим ширине импульса, длине записи, решётке, белому гауссовскому шуму и направляющему вектору без задержек. Изменение этих условий требует отдельного исследования. Физическая измерительная калибровка и модель окрашенного шума в данном эксперименте не рассматривались.')

h('Приложение. Воспроизводимость и сохранность исходного кода')
p('Рабочий ноутбук lorentz_research.ipynb содержит выполненные ячейки, графики и числовые результаты. Ячейки исходного notebook.ipynb перенесены с добавлением сохранения рисунков; исходные функции сигнала, энергии, генерации каналов и оценивания использованы без изменений. Статистическая схема report.ipynb развёрнута через list_errors для сохранения отдельных ошибок и вычисления дополнительных статистик; её дисперсии проверены на совпадение с var_vs_params в двух контрольных точках.')
p('Все исходные ноутбуки, файлы py_files/*.py, tests/*.py и оба исходных документа проверены по SHA-256 до и после расчёта. Контрольные суммы совпали; пользовательские изменения, существовавшие до работы, сохранены. Манифест проверки находится в source_integrity.json. Временного редактирования исходных файлов не потребовалось.')
if (OUT/'source_postcheck.json').exists():
    postcheck=json.loads((OUT/'source_postcheck.json').read_text(encoding='utf-8'))
    if postcheck['changed_after_calculation']:
        p('После завершения расчёта при отдельной финальной проверке обнаружены внешние изменения notebook.ipynb и py_files/functions.py. Они не откатывались. Контрольные суммы этой проверки записаны в source_postcheck.json; числовые результаты отчёта относятся к версии, зафиксированной в source_integrity.json. Для воспроизведения уже подготовленных рисунков их данные и выполненные ячейки сохранены в рабочем ноутбуке.')
p(f'Среда: Python 3.13.5, NumPy {r["versions"]["numpy"]}, Matplotlib {r["versions"]["matplotlib"]}. Повторный расчёт выполняет run_study.py, сборку отчёта — build_report.py; команды приведены в README.md. Для сборки используются python-docx, latex2mathml и преобразование MathML → OMML из Microsoft Office. Формулы DOCX редактируются в Word.')
p('Файлы результатов: results.json — сводные числа; statistics.csv — полная таблица статистик; errors.json — индивидуальные ошибки; comparison_intervals.json — интервалы отношения дисперсий; figures/ — рисунки PNG.')
dest=OUT/'lorentz_research_report_final.docx'
doc.save(dest)
(OUT/'lorentz_research_report_final.md').write_text('\n'.join(md),encoding='utf-8')
print(dest)
print('Equations:',eqno,'Figures:',figno,'Bootstrap:',intervals)
