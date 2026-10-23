"""Build a concise report with full-size, displayed Word equations."""
from pathlib import Path
import json
import math
import sys

sys.dont_write_bytecode = True
OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
sys.path.insert(0, str(ROOT))

import numpy as np
import matplotlib
matplotlib.use('Agg')
from matplotlib import pyplot as plt
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from latex2mathml.converter import convert
from py_files.functions import get_impulse_signal, calculate_freq_spectr, calculate_energy

results = json.loads((OUT / 'results.json').read_text(encoding='utf-8'))
checks = results['checks']
params = results['parameters']
FIG = OUT / 'figures'

# Reuse the original signal/FFT functions with the saved experiment parameters.
# The source notebooks and modules are never written by this script.
t = np.arange(params['N']) / params['fs']
signal_config = dict(A=checks['A'], duration_pulse=checks['tau'],
                     center_pulse=checks['tc'], time_vector=t)
pulse = get_impulse_signal(signal_config)
assert np.isclose(calculate_energy(pulse, t), checks['E_imp_numeric'], rtol=1e-12)
frequencies = np.fft.fftshift(np.fft.fftfreq(params['N'], 1 / params['fs']))
spectrum = calculate_freq_spectr(pulse, params['fs'])
theory = np.pi * checks['A'] * np.exp(-2*np.pi*checks['tau']*np.abs(frequencies))

plt.rcParams.update({'font.family':'DejaVu Sans', 'font.size':11,
                     'axes.labelsize':12, 'legend.fontsize':10,
                     'axes.titlesize':12, 'lines.linewidth':1.7})

fig, ax = plt.subplots(figsize=(6.45, 2.5), layout='constrained')
ax.plot(t, pulse, color='#2056a0')
ax.axvline(checks['tc'], color='black', ls='--', lw=1, label=r'$t_c=T/2$')
ax.set(xlabel='Время, с', ylabel='Давление, Па', xlim=(0, checks['T']))
ax.grid(alpha=.3); ax.legend()
fig.savefig(FIG / 'short_pulse.png', dpi=220)
plt.close(fig)

fig, ax = plt.subplots(figsize=(6.45, 2.5), layout='constrained')
ax.plot(frequencies, spectrum, label='ДПФ конечной записи')
ax.plot(frequencies, theory, '--', label='Аналитический спектр')
ax.set(xlabel='Частота, Гц', ylabel=r'$|\widetilde P(f)|$, Па·с', xlim=(-25,25))
ax.grid(alpha=.3); ax.legend()
fig.savefig(FIG / 'short_spectrum.png', dpi=220)
plt.close(fig)

fig, ax = plt.subplots(figsize=(6.45, 2.8), layout='constrained')
for signal_name, linestyle in [('Лоренц','-'),('Синусоида','--')]:
    for method, color in [('Бартлетт','#2056a0'),('Кейпон','#bd3a32')]:
        rows = [d for d in results['statistics'] if d['signal']==signal_name and d['method']==method]
        ax.plot([d['snr'] for d in rows], np.degrees([d['rmse_rad'] for d in rows]),
                linestyle, color=color, marker='.', label=f'{signal_name}: {method}')
ax.set(xlabel='SNR, линейная шкала', ylabel='RMSE, град.')
ax.grid(alpha=.3); ax.legend()
fig.savefig(FIG / 'short_comparison.png', dpi=220)
plt.close(fig)

doc = Document(ROOT / 'reports/dox/info_snr_supplement.docx')
for child in list(doc._element.body):
    if child.tag != qn('w:sectPr'):
        doc._element.body.remove(child)
section = doc.sections[0]
section.page_width=Cm(21); section.page_height=Cm(29.7)
section.left_margin=Cm(3); section.right_margin=Cm(1.5)
section.top_margin=Cm(2); section.bottom_margin=Cm(2)

for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    s=doc.styles[name]
    s.font.name='Times New Roman'; s.font.size=Pt(14)
    s.font.color.rgb=RGBColor(0,0,0)
    # Remove theme-font overrides inherited from the supplied Word template.
    fonts=s.element.get_or_add_rPr().find(qn('w:rFonts'))
    for key in list(fonts.attrib):
        if 'Theme' in key or 'theme' in key:
            del fonts.attrib[key]
    s.paragraph_format.space_before=Pt(0)
    s.paragraph_format.space_after=Pt(0)
    s.paragraph_format.line_spacing=1.3
    s.paragraph_format.first_line_indent=Cm(0)

normal=doc.styles['Normal'].paragraph_format
normal.first_line_indent=Cm(1.25)
normal.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name in ['Heading 1','Heading 2']:
    doc.styles[name].font.bold=True
    doc.styles[name].paragraph_format.space_before=Pt(8)
    doc.styles[name].paragraph_format.space_after=Pt(3)
    doc.styles[name].paragraph_format.keep_with_next=True

xsl=Path('C:/Program Files/Microsoft Office/Office16/MML2OMML.XSL')
transform=etree.XSLT(etree.parse(str(xsl)))
markdown=[]
eq_count=0
fig_count=0

def set_font_properties(properties):
    fonts=OxmlElement('w:rFonts')
    fonts.set(qn('w:ascii'),'Cambria Math'); fonts.set(qn('w:hAnsi'),'Cambria Math')
    properties.append(fonts)
    size=OxmlElement('w:sz'); size.set(qn('w:val'),'28'); properties.append(size)
    size_cs=OxmlElement('w:szCs'); size_cs.set(qn('w:val'),'28'); properties.append(size_cs)

def p(text):
    pp=doc.add_paragraph(text)
    pp.paragraph_format.widow_control=True
    markdown.append(text+'\n')
    return pp

def h(text):
    doc.add_heading(text, level=2)
    markdown.append('## '+text+'\n')

def eq(tex):
    """Display math, not inline math: full-size fractions and integral limits."""
    global eq_count
    eq_count+=1
    pp=doc.add_paragraph()
    pp.paragraph_format.first_line_indent=Cm(0)
    pp.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pp.paragraph_format.line_spacing=1
    pp.paragraph_format.space_before=Pt(5)
    pp.paragraph_format.space_after=Pt(5)
    pp.paragraph_format.keep_together=True
    paragraph_rpr=OxmlElement('w:rPr')
    set_font_properties(paragraph_rpr)
    pp._p.get_or_add_pPr().append(paragraph_rpr)
    mathml=etree.fromstring(convert(tex, display='block').encode())
    omath=transform(mathml).getroot()
    for mr in omath.iter(qn('m:r')):
        properties=OxmlElement('w:rPr')
        set_font_properties(properties)
        mpr=mr.find(qn('m:rPr'))
        if mpr is None:
            mr.insert(0,properties)
        else:
            mr.insert(list(mr).index(mpr)+1,properties)
    for tag in ['fPr','naryPr','radPr','sSubPr','sSupPr','sSubSupPr','dPr']:
        for pr in omath.iter(qn('m:'+tag)):
            ctrl=OxmlElement('m:ctrlPr')
            rpr=OxmlElement('w:rPr'); set_font_properties(rpr)
            ctrl.append(rpr); pr.append(ctrl)
    display=OxmlElement('m:oMathPara')
    display_pr=OxmlElement('m:oMathParaPr')
    jc=OxmlElement('m:jc'); jc.set(qn('m:val'),'center')
    display_pr.append(jc); display.append(display_pr)
    display.append(omath)
    pp._p.append(display)
    markdown.append('$$\n'+tex+'\n$$\n')

def figure(filename,caption):
    global fig_count
    fig_count+=1
    pp=doc.add_paragraph()
    pp.paragraph_format.first_line_indent=Cm(0)
    pp.paragraph_format.keep_with_next=True
    pp.paragraph_format.line_spacing=1
    pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pp.add_run().add_picture(str(FIG/filename),width=Cm(16.4))
    pp=doc.add_paragraph(f'Рисунок {fig_count} — {caption}',style='Caption')
    pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pp.paragraph_format.line_spacing=1
    pp.paragraph_format.keep_together=True
    pp.paragraph_format.space_after=Pt(4)
    for run in pp.runs:
        run.font.size=Pt(12); run.italic=False
    markdown.append(f'![{caption}](figures/{filename})\n')

def table(headers,rows):
    tt=doc.add_table(rows=1, cols=len(headers)); tt.style='Table Grid'
    for cell,text in zip(tt.rows[0].cells,headers): cell.text=str(text)
    for row in rows:
        for cell,text in zip(tt.add_row().cells,row):cell.text=str(text)
    for i,row in enumerate(tt.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for cell in row.cells:
            for pp in cell.paragraphs:
                pp.paragraph_format.first_line_indent=Cm(0)
                pp.paragraph_format.line_spacing=1.1
                pp.paragraph_format.space_before=Pt(4)
                pp.paragraph_format.space_after=Pt(4)
                pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in pp.runs:
                    run.font.name='Times New Roman';run.font.size=Pt(12);run.bold=i==0
    markdown.append('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+''.join('| '+' | '.join(row)+' |\n' for row in rows)+'\n')

def num(v,places=4): return f'{v:.{places}f}'.replace('.',',')

pp=doc.add_paragraph('ИССЛЕДОВАНИЕ ИМПУЛЬСА ЛОРЕНЦА',style='Title')
pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
pp.paragraph_format.space_after=Pt(6)
for run in pp.runs:run.bold=True
markdown.append('# Исследование импульса Лоренца\n')

p('Параметры: M = 3 АВП; R = 1,7163 м; c = 343,26 м/с; fₛ = 8800 Гц; N = 10000; τ = 0,05 с; истинный азимут φ₀ = 45°. Синусоида сравнения: амплитуда 1 Па, частота 10 Гц.')

h('1 Отношение сигнал/шум')
p('Отношение SNR задаётся через энергию сигнала и математическое ожидание энергии шума на одном интервале наблюдения:')
eq(r'\mathrm{SNR}=\frac{E_s(T)}{\mathbb{E}[E_n(T)]}=\frac{\displaystyle\int_0^T P^2(t)\,dt}{T\sigma^2}')
p('Здесь σ² — дисперсия гауссовских шумовых отсчётов. Для принятого расчёта методом трапеций средняя энергия шума равна Tσ². Длительность записи определяется первым и последним отсчётами:')
eq(r'T=\frac{N-1}{f_s}=1{,}13625\ \mathrm{с}')
p('Интегрируем от 0 до T, поскольку алгоритм обрабатывает только эту запись. Хвосты импульса вне окна в расчёт SNR не входят. Уровень шума задаётся по фактической энергии в окне:')
eq(r'\sigma=\sqrt{\frac{E_s(T)}{T\,\mathrm{SNR}}}')

h('2 Импульс и положение его центра')
eq(r'P(t)=\frac{A\tau}{\tau^2+(t-t_c)^2},\qquad t_c=\frac{T}{2}')
p('В расчёте A = 1 Па·с, поэтому высота пика A/τ = 20 Па. Параметр τ — половина ширины на половине максимума; полная ширина равна 0,1 с. Центр tс = 0,568125 с.')
p('Положение T/2 даёт одинаковый запас до границ окна. Для симметричного импульса, убывающего от центра, оно максимизирует энергию в записи. При заданных параметрах в окно попадает 99,9713 % полной энергии. На m-м АВП сигнал P(t − dₘ) имеет центр tс + dₘ, где dₘ — задержка распространения.')

doc.add_page_break()
figure('short_pulse.png','Импульс Лоренца с центром T/2')
h('3 Спектр импульса')
p('Для преобразования Фурье по частоте f в герцах:')
eq(r'\widetilde P(f)=\int_{-\infty}^{\infty}P(t)e^{-i2\pi ft}\,dt')
p('Спектр полной импульсной формы и его модуль равны:')
eq(r'\widetilde P_\infty(f)=\pi A e^{-2\pi\tau|f|}e^{-i2\pi ft_c}')
eq(r'|\widetilde P_\infty(f)|=\pi|A|e^{-2\pi\tau|f|}')
p('Сдвиг центра меняет фазу, сохраняя модуль спектра полной формы. Максимум расположен при f = 0; с ростом |f| спектр экспоненциально убывает.')
figure('short_spectrum.png','Модуль спектра: ДПФ записи и аналитическая форма')
p('ДПФ нормирована множителем 1/fₛ. Отличие около нуля вызвано конечным окном: ДПФ даёт 2,9660 Па·с, а полная форма — πA ≈ 3,1416 Па·с.')

doc.add_page_break()
h('4 Оценка азимута и сравнение с синусоидой')
p('Методы Бартлетта и Кейпона сравнивались при одинаковом энергетическом SNR, одной решётке и одной записи. Для каждого сигнала шум нормировался по его энергии. На каждую точку SNR выполнено 100 реализаций; seed = 10.')
p('Точность характеризуется корнем средней квадратичной ошибки:')
eq(r'\mathrm{RMSE}=\sqrt{\frac{1}{K}\sum_{k=1}^{K}\varepsilon_k^2},\qquad K=100')
p('Здесь εₖ — круговая разность оценки и истинного азимута. В таблице приведены оценка для одной реализации при SNR = 2 и RMSE по 100 реализациям при том же SNR.')
rows=[]
for single in results['single']:
    st=next(d for d in results['statistics'] if d['signal']==single['signal'] and d['method']==single['method'] and d['snr']==2)
    rows.append((single['signal'],single['method'],num(single['estimate_deg']),num(math.degrees(st['rmse_rad']))))
table(['Сигнал','Метод','Оценка, град.','RMSE, град.'],rows)
figure('short_comparison.png','Ошибка оценки азимута в зависимости от SNR')
p('С ростом SNR ошибки в целом уменьшаются. При SNR = 2 RMSE составляет около 0,24° для импульса и 0,26° для синусоиды. Точность методов близка. Небольшая разница между сигналами в этой выборке не доказывает устойчивого преимущества импульса.')

dest=OUT/'lorentz_research_report_short.docx'
doc.save(dest)
(OUT/'lorentz_research_report_short.md').write_text('\n'.join(markdown),encoding='utf-8')
print(dest)
print('Displayed 14 pt equations:',eq_count,'Figures:',fig_count)
